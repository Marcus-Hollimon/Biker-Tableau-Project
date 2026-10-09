import openpyxl, sys
src, dst = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src)
a, b = wb['year 2011'], wb['year 2012']
key = lambda row: (row[0].replace(year=2011),) + tuple(row[1:])
dupes = {key(r) for r in b.iter_rows(min_row=2, values_only=True)}
rows = list(a.iter_rows(min_row=2, values_only=True))
keep = [r for r in rows if key(r) not in dupes]
print('removed', len(rows) - len(keep), 'kept', len(keep))
a.delete_rows(2, a.max_row)
for r in keep:
    a.append(r)
    a.cell(a.max_row, 1).number_format = 'm/d/yy h:mm'
wb.save(dst)
